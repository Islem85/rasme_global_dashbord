"""
Exporte la liste des sites NON cartographiables depuis ``public.kobo_clean``
vers ``assets/sites_sans_gps.csv`` et ``assets/sites_sans_gps.xlsx`` (mis en
forme).

Définition « sans GPS » — identique au compteur du dashboard (global_view.py) :
un site au statut valide qui n'est pas affichable sur la carte, c.-à-d. dont
les coordonnées sont :
  • absentes (latitude/longitude NULL)               → « Coordonnées manquantes »
  • présentes mais hors de la bounding box Afrique    → « Coordonnées hors
    Afrique (erreur de saisie) »

La bbox et le filtre de statut sont partagés avec db.map_points via config, de
sorte que le compte de ce fichier == le « N site(s) sans coordonnées GPS »
montré sous la carte.

Lancé en dernière étape de run_refresh_all.py (après cleanup_kobo.sql qui
construit kobo_clean). Peut aussi être exécuté seul :  python export_sites_sans_gps.py
"""

from __future__ import annotations

import logging
import os
import sys
from datetime import date
from pathlib import Path

import pandas as pd
import psycopg2
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import (
    AFRICA_LAT_MAX,
    AFRICA_LAT_MIN,
    AFRICA_LON_MAX,
    AFRICA_LON_MIN,
    COL,
    INVALID_STATUS,
    country_label,
)

HERE = Path(__file__).resolve().parent
load_dotenv(HERE / ".env")

PG_HOST     = os.getenv("PG_HOST", "localhost")
PG_PORT     = int(os.getenv("PG_PORT", "5432"))
PG_DB       = os.getenv("PG_DB", "mapping")
PG_USER     = os.getenv("PG_USER", "postgres")
PG_PASSWORD = os.getenv("PG_PASSWORD", "")

ASSETS   = HERE / "assets"
CSV_PATH = ASSETS / "sites_sans_gps.csv"
XLSX_PATH = ASSETS / "sites_sans_gps.xlsx"

(HERE / "logs").mkdir(exist_ok=True)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(HERE / "logs" / "export_sites_sans_gps.log", encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)

# Palette AfDB / RASME (cohérente avec config.PALETTE).
GREEN_DK = "005C33"
GREEN    = "008F4F"
RED_BG   = "FCE4E2"
FONT     = "Arial"

# Ordre + en-têtes FR des colonnes exportées. La clé est l'alias SQL (snake).
OUT_COLS = {
    "country":                  "Pays",
    "nom_projet":               "Projet",
    "nom_site":                 "Site",
    "secteur_activite":         "Secteur",
    "statut_implementation":    "Statut",
    "niveau_realisation":       "Niveau de réalisation",
    "description_activite":     "Description",
    "region":                   "Région",
    "province":                 "Province",
    "commune":                  "Commune",
    "date_collecte":            "Date collecte",
    "date_debut":               "Date début",
    "date_fin_prevue":          "Date fin prévue",
    "nom_agence_execution":     "Agence d'exécution",
    "institutions_financement": "Financement",
    "raison":                   "Raison",
    "latitude":                 "Latitude (brut)",
    "longitude":                "Longitude (brut)",
    "identifiant_pa":           "Identifiant PA",
    "submission_uuid":          "UUID soumission",
}

# Largeurs de colonnes Excel (par en-tête FR).
WIDTHS = {
    "Pays": 16, "Projet": 42, "Site": 30, "Secteur": 18, "Statut": 13,
    "Niveau de réalisation": 16, "Description": 50, "Région": 12, "Province": 14,
    "Commune": 16, "Date collecte": 18, "Date début": 18, "Date fin prévue": 18,
    "Agence d'exécution": 28, "Financement": 26, "Raison": 34,
    "Latitude (brut)": 14, "Longitude (brut)": 14, "Identifiant PA": 22,
    "UUID soumission": 38,
}
WRAP_COLS = {"Projet", "Site", "Description"}


def fetch() -> pd.DataFrame:
    """Lit kobo_clean et renvoie les sites non cartographiables (libellés FR)."""
    invalid = ", ".join("'" + s.replace("'", "''") + "'" for s in INVALID_STATUS)
    lat, lon = COL["lat"], COL["lon"]
    status   = COL["status"]
    status_ok  = f"{status} IS NOT NULL AND {status} NOT IN ({invalid})"
    has_coords = f"{lat} IS NOT NULL AND {lon} IS NOT NULL"
    in_africa  = (f"{lon} BETWEEN {AFRICA_LON_MIN} AND {AFRICA_LON_MAX} "
                  f"AND {lat} BETWEEN {AFRICA_LAT_MIN} AND {AFRICA_LAT_MAX}")

    sql = f"""
        SELECT
            {COL['country']}             AS country,
            {COL['project_title']}       AS nom_projet,
            {COL['site_name']}           AS nom_site,
            {COL['sector']}              AS secteur_activite,
            {COL['status']}              AS statut_implementation,
            {COL['progress_band']}       AS niveau_realisation,
            {COL['site_description']}    AS description_activite,
            {COL['region_name']}         AS region,
            {COL['district_name']}       AS province,
            {COL['commune']}             AS commune,
            {COL['collection_date']}     AS date_collecte,
            {COL['planned_start_date']}  AS date_debut,
            {COL['planned_end_date']}    AS date_fin_prevue,
            {COL['implementing_agency']} AS nom_agence_execution,
            {COL['funder']}              AS institutions_financement,
            CASE WHEN NOT ({has_coords}) THEN 'Coordonnées manquantes'
                 ELSE 'Coordonnées hors Afrique (erreur de saisie)' END AS raison,
            {COL['lat']}                 AS latitude,
            {COL['lon']}                 AS longitude,
            {COL['project_code']}        AS identifiant_pa,
            submission_uuid
        FROM public.kobo_clean
        WHERE {status_ok}
          AND NOT ({has_coords} AND {in_africa})
        ORDER BY country, nom_projet, nom_site
    """
    conn = psycopg2.connect(
        host=PG_HOST, port=PG_PORT, dbname=PG_DB,
        user=PG_USER, password=PG_PASSWORD,
    )
    try:
        with conn.cursor() as cur:
            cur.execute(sql)
            rows = cur.fetchall()
            cols = [d[0] for d in cur.description]
    finally:
        conn.close()

    df = pd.DataFrame(rows, columns=cols)
    # slug pays -> libellé français.
    df["country"] = df["country"].map(lambda s: country_label(s, "FR") if s else s)
    return df.rename(columns=OUT_COLS)[list(OUT_COLS.values())]


def write_xlsx(df: pd.DataFrame) -> None:
    """Classeur mis en forme : feuille Synthèse + feuille de données filtrable."""
    cols = list(df.columns)
    n = len(df)
    n_manq = int((df["Raison"] == "Coordonnées manquantes").sum())
    n_hors = int(df["Raison"].str.contains("hors Afrique", regex=False).sum())
    pays_counts = df["Pays"].value_counts()

    wb = Workbook()

    # ── Feuille données ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Sites sans GPS"
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GREEN_DK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, (_, row) in enumerate(df.iterrows(), start=2):
        for c, name in enumerate(cols, start=1):
            val = row[name]
            cell = ws.cell(row=r, column=c, value=(None if pd.isna(val) else val))
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(name in WRAP_COLS))

    ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
    table = Table(displayName="SitesSansGPS", ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium7", showRowStripes=True, showColumnStripes=False,
        showFirstColumn=False, showLastColumn=False,
    )
    ws.add_table(table)
    for c, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS.get(name, 16)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    raison_col = get_column_letter(cols.index("Raison") + 1)
    ws.conditional_formatting.add(
        f"{raison_col}2:{raison_col}{n + 1}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("hors Afrique",${raison_col}2))'],
                    fill=PatternFill("solid", fgColor=RED_BG)),
    )

    # ── Feuille synthèse ─────────────────────────────────────────────
    syn = wb.create_sheet("Synthèse", 0)
    syn.sheet_view.showGridLines = False
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    syn.merge_cells("A1:C1")
    t = syn["A1"]
    t.value = "Sites sans coordonnées GPS — RASME / BAD"
    t.font = Font(name=FONT, bold=True, size=16, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=GREEN)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    syn.row_dimensions[1].height = 34
    syn.merge_cells("A2:C2")
    syn["A2"] = f"Généré le {date.today():%d/%m/%Y} · source : public.kobo_clean"
    syn["A2"].font = Font(name=FONT, size=9, italic=True, color="5C6770")

    def hdr(cell):
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GREEN_DK)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.border = border
        return cell

    def label(cell, bold=False):
        cell.font = Font(name=FONT, size=10, bold=bold)
        cell.alignment = Alignment(horizontal="left", indent=1)
        cell.border = border
        return cell

    def num(cell, bold=False):
        cell.font = Font(name=FONT, size=10, bold=bold)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
        return cell

    r = 4
    hdr(syn.cell(r, 1, "Par raison")); hdr(syn.cell(r, 2, "Sites"))
    syn.cell(r, 3).border = border
    r += 1; label(syn.cell(r, 1, "Coordonnées manquantes (NULL)")); num(syn.cell(r, 2, n_manq))
    r += 1; label(syn.cell(r, 1, "Coordonnées hors Afrique (erreur)")); num(syn.cell(r, 2, n_hors))
    r += 1; label(syn.cell(r, 1, "Total"), bold=True); num(syn.cell(r, 2, n), bold=True)

    r += 2
    hdr(syn.cell(r, 1, "Par pays")); hdr(syn.cell(r, 2, "Sites"))
    syn.cell(r, 3).border = border
    for pays in pays_counts.index.tolist():
        r += 1
        label(syn.cell(r, 1, pays))
        num(syn.cell(r, 2, int(pays_counts[pays])))

    syn.column_dimensions["A"].width = 36
    syn.column_dimensions["B"].width = 10
    syn.column_dimensions["C"].width = 3

    wb.save(XLSX_PATH)


def main() -> int:
    log.info("export_sites_sans_gps : démarrage")
    ASSETS.mkdir(exist_ok=True)
    try:
        df = fetch()
    except Exception as exc:
        log.error(f"ERREUR lecture kobo_clean : {type(exc).__name__}: {exc}")
        return 1

    n = len(df)
    df.to_csv(CSV_PATH, index=False, sep=";", encoding="utf-8-sig")
    write_xlsx(df)

    n_manq = int((df["Raison"] == "Coordonnées manquantes").sum())
    n_hors = n - n_manq
    log.info(f"{n} sites sans GPS  ({n_manq} coordonnées manquantes, "
             f"{n_hors} hors Afrique)")
    log.info(f"  -> {CSV_PATH.name} + {XLSX_PATH.name} écrits dans assets/")
    return 0


if __name__ == "__main__":
    sys.exit(main())
