"""
Exporte la liste des « Projets cartographiés sans correspondance dans la base
BAD » vers ``assets/projets_sans_bad.csv`` et ``assets/projets_sans_bad.xlsx``
(mis en forme).

Définition — identique au KPI du dashboard (global_view.py / country_view.py,
explainer « Projets cartographiés sans correspondance dans la base BAD ») :
un projet présent dans les données terrain (kobo_clean, donc *cartographié*)
dont le code SAP n'existe dans AUCUN enregistrement de la table projet_BAD des
pays RASME. C'est le compteur ``coverage_overview()['n_internal_only']`` et,
au détail, les lignes ``coverage_status == 'Non cartographié'`` de
``coverage_table(None)``.

La logique de rapprochement (code court identifiant_pa → 4 segments, ou
nom_projet → public."code_SAP") est REUTILISÉE depuis db.py pour garantir un
compte identique à celui affiché. On suit le pattern de smoke_test.py : on
neutralise le cache Streamlit avant d'importer db, ce qui rend les fonctions
@st.cache_data appelables hors runtime Streamlit.

Lancé par run_refresh_all.py après cleanup_kobo.sql + scrape_mapafrica.py
(il a besoin de kobo_clean ET de projet_BAD). Exécutable seul :
    python export_projets_sans_bad.py
"""

from __future__ import annotations

import logging
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
load_dotenv(HERE / ".env")


# ── Neutralise le cache Streamlit pour importer db hors runtime (cf. smoke_test) ──
class _Stub:
    def __getattr__(self, name):
        def passthrough(*a, **k):
            if a and callable(a[0]):
                return a[0]
            def deco(fn):
                return fn
            return deco
        return passthrough


import streamlit  # noqa: E402

streamlit.cache_data = _Stub().__getattr__("x")
streamlit.cache_resource = _Stub().__getattr__("x")

import db  # noqa: E402
from config import country_label  # noqa: E402

ASSETS    = HERE / "assets"
CSV_PATH  = ASSETS / "projets_sans_bad.csv"
XLSX_PATH = ASSETS / "projets_sans_bad.xlsx"

(HERE / "logs").mkdir(exist_ok=True)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(HERE / "logs" / "export_projets_sans_bad.log", encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
# Les warnings "missing ScriptRunContext" de Streamlit hors runtime ne sont pas
# pertinents ici.
logging.getLogger("streamlit").setLevel(logging.ERROR)
log = logging.getLogger(__name__)

GREEN_DK = "005C33"
GREEN    = "008F4F"
FONT     = "Arial"

# En-têtes FR (clé = colonne interne).
OUT_COLS = {
    "code":              "Code projet (SAP)",
    "project_title":     "Projet",
    "all_titles":        "Autres intitulés",
    "pays":              "Pays",
    "nb_sites":          "Nb sites cartographiés",
    "secteurs":          "Secteur(s)",
    "statuts":           "Statut(s)",
    "financements":      "Financement(s)",
    "regions":           "Région(s)",
    "derniere_collecte": "Dernière collecte",
}
WIDTHS = {
    "Code projet (SAP)": 24, "Projet": 48, "Autres intitulés": 40, "Pays": 18,
    "Nb sites cartographiés": 13, "Secteur(s)": 22, "Statut(s)": 22,
    "Financement(s)": 34, "Région(s)": 18, "Dernière collecte": 16,
}
WRAP_COLS = {"Projet", "Autres intitulés", "Financement(s)", "Secteur(s)", "Statut(s)"}

# Agrégats kobo_clean par code SAP résolu (mêmes deux chemins que db.py).
ENRICH_SQL = """
WITH rc AS (
    SELECT c.code,
           NULLIF(btrim(x.country), '')                  AS country,
           NULLIF(btrim(x.secteur_activite), '')         AS secteur,
           NULLIF(btrim(x.statut_implementation), '')    AS statut,
           NULLIF(btrim(x.institutions_financement), '') AS financement,
           NULLIF(btrim(x.region), '')                   AS region,
           CASE WHEN x.date_collecte ~ '^[0-9]{4}-[0-9]{2}-[0-9]{2}'
                THEN left(x.date_collecte, 10) END        AS d,
           x.submission_uuid
    FROM (
        SELECT k.country, k.secteur_activite, k.statut_implementation,
               k.institutions_financement, k.region, k.date_collecte,
               k.submission_uuid,
               btrim(array_to_string((string_to_array(k.identifiant_pa, '-'))[1:4], '-')) AS code_a,
               (SELECT btrim(s."Code_SAP") FROM public."code_SAP" s
                 WHERE btrim(s."Name of the project") = btrim(k.nom_projet) LIMIT 1)      AS code_b
        FROM public.kobo_clean k
        WHERE k.statut_implementation IS NOT NULL
          AND k.statut_implementation NOT IN ('#status', 'None', '')
    ) x
    CROSS JOIN LATERAL (VALUES (x.code_a), (x.code_b)) AS c(code)
    WHERE c.code IS NOT NULL AND c.code <> ''
)
SELECT code,
       count(DISTINCT submission_uuid)               AS nb_sites,
       string_agg(DISTINCT country, ', ')            AS pays_slugs,
       string_agg(DISTINCT secteur, ', ')            AS secteurs,
       string_agg(DISTINCT statut, ', ')             AS statuts,
       string_agg(DISTINCT financement, ', ')        AS financements,
       string_agg(DISTINCT region, ', ')             AS regions,
       max(d)                                        AS derniere_collecte
FROM rc
GROUP BY code
"""


def _label_pays(slugs: str | None) -> str:
    if not slugs:
        return ""
    parts = [p.strip() for p in str(slugs).split(",") if p.strip()]
    return ", ".join(country_label(p, "FR") for p in parts)


def build() -> pd.DataFrame:
    """Renvoie le DataFrame final (libellés FR, trié) des projets sans BAD."""
    cov = db.coverage_table(None)
    unm = cov[cov["coverage_status"] == "Non cartographié"][
        ["code", "project_title", "all_titles"]
    ].copy()

    enrich = db._q(ENRICH_SQL)
    df = unm.merge(enrich, on="code", how="left")

    df["pays"] = df["pays_slugs"].map(_label_pays)
    df["nb_sites"] = df["nb_sites"].fillna(0).astype(int)
    df = df.rename(columns=OUT_COLS)[list(OUT_COLS.values())]
    df = df.sort_values(
        ["Nb sites cartographiés", "Code projet (SAP)"], ascending=[False, True]
    ).reset_index(drop=True)
    return df


def write_xlsx(df: pd.DataFrame) -> None:
    cols = list(df.columns)
    n = len(df)
    total_sites = int(df["Nb sites cartographiés"].sum())
    pays_counts = (
        df["Pays"].fillna("—").str.split(", ").explode().str.strip()
        .replace("", "—").value_counts()
    )

    wb = Workbook()

    # ── Feuille données ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Projets sans BAD"
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GREEN_DK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, (_, row) in enumerate(df.iterrows(), start=2):
        for c, name in enumerate(cols, start=1):
            val = row[name]
            cell = ws.cell(row=r, column=c, value=(None if pd.isna(val) else val))
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(name in WRAP_COLS))

    ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
    table = Table(displayName="ProjetsSansBAD", ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium7", showRowStripes=True, showColumnStripes=False,
        showFirstColumn=False, showLastColumn=False,
    )
    ws.add_table(table)
    for c, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS.get(name, 16)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # ── Feuille synthèse ─────────────────────────────────────────────
    syn = wb.create_sheet("Synthèse", 0)
    syn.sheet_view.showGridLines = False
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    syn.merge_cells("A1:C1")
    t = syn["A1"]
    t.value = "Projets cartographiés sans correspondance BAD — RASME"
    t.font = Font(name=FONT, bold=True, size=14, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=GREEN)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    syn.row_dimensions[1].height = 34
    syn.merge_cells("A2:C2")
    syn["A2"] = (f"Généré le {date.today():%d/%m/%Y} · source : public.kobo_clean × "
                 'public."projet_BAD"')
    syn["A2"].font = Font(name=FONT, size=9, italic=True, color="5C6770")

    def hdr(cell):
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GREEN_DK)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.border = border
        return cell

    def label(cell, bold=False):
        cell.font = Font(name=FONT, size=10, bold=bold)
        cell.alignment = Alignment(horizontal="left", indent=1)
        cell.border = border
        return cell

    def num(cell, bold=False):
        cell.font = Font(name=FONT, size=10, bold=bold)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
        return cell

    r = 4
    hdr(syn.cell(r, 1, "Indicateur")); hdr(syn.cell(r, 2, "Valeur"))
    syn.cell(r, 3).border = border
    r += 1; label(syn.cell(r, 1, "Projets sans correspondance BAD"), bold=True); num(syn.cell(r, 2, n), bold=True)
    r += 1; label(syn.cell(r, 1, "Sites cartographiés concernés")); num(syn.cell(r, 2, total_sites))

    r += 2
    hdr(syn.cell(r, 1, "Par pays")); hdr(syn.cell(r, 2, "Projets"))
    syn.cell(r, 3).border = border
    for pays in pays_counts.index.tolist():
        r += 1
        label(syn.cell(r, 1, pays))
        num(syn.cell(r, 2, int(pays_counts[pays])))
    r += 1
    note = syn.cell(r, 1, "(un projet multinational est compté dans chaque pays concerné)")
    note.font = Font(name=FONT, size=8, italic=True, color="5C6770")
    syn.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

    syn.column_dimensions["A"].width = 38
    syn.column_dimensions["B"].width = 10
    syn.column_dimensions["C"].width = 3

    wb.save(XLSX_PATH)


def main() -> int:
    log.info("export_projets_sans_bad : démarrage")
    ASSETS.mkdir(exist_ok=True)
    try:
        df = build()
        expected = db.coverage_overview().get("n_internal_only")
    except Exception as exc:
        log.error(f"ERREUR : {type(exc).__name__}: {exc}")
        return 1

    n = len(df)
    if expected is not None and n != int(expected):
        log.warning(f"Incohérence : {n} lignes exportées vs n_internal_only={expected} "
                    "(le détail par code et le compteur agrégé diffèrent).")

    df.to_csv(CSV_PATH, index=False, sep=";", encoding="utf-8-sig")
    write_xlsx(df)
    log.info(f"{n} projets cartographiés sans correspondance BAD "
             f"({int(df['Nb sites cartographiés'].sum())} sites concernés)")
    log.info(f"  -> {CSV_PATH.name} + {XLSX_PATH.name} écrits dans assets/")
    return 0


if __name__ == "__main__":
    sys.exit(main())
